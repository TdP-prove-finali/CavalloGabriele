from openpyxl import load_workbook

from dao.data_import.balance_dao import RawCompanyBalanceImporter
from dao.data_import.comp_preproc_dao import CompanyPreprocessorDAO
from model.raw_company_data import CompanyData, RawCompanyData


class CompanyImportDAO:
    def __init__(self, filePath):
        self.__filePath__ = filePath
        self.__company__: CompanyData | None = None
        self.__translations__: dict = None

    def import_single_company(self):
        """
        Importa una company a partire da un file Excel eseguendo tutti gli step, ossia:
        1 - Importazione della struttura Raw dal file excel
        2 - Pre processamento della struttura per rendela di facile manipolazione

        Successivamente utilizzare {self.company} e {self.translations} per accedere ai dati della company e alle traduzioni
        automatiche di tutti i campi con i rispettivi descrittori completi.
        :return:
        """
        wb = load_workbook(filename=self.__filePath__, read_only=False, data_only=True)

        if len(wb.sheetnames) == 0:
            raise Exception("Il file non ha fogli da cui caricare i dati")

        firstSheet = wb.worksheets[0]

        dao = RawCompanyBalanceImporter(firstSheet)
        dao.import_from_excel()
        raw_company = dao.company

        self.__company__ = CompanyPreprocessorDAO.process_company_data(raw_company)
        self.__translations__ = dao.get_translations()

    def import_company_chunk(self) -> list[CompanyData]:
        imported = []
        """
        Importa tutte le company presenti nel file excel, una per foglio di lavoro
        :return: 
        """
        wb = load_workbook(filename=self.__filePath__, read_only=False, data_only=True)

        if len(wb.sheetnames) == 0:
            raise Exception("Il file non ha fogli da cui caricare i dati")

        for sheet in wb.worksheets:
            try:
                imp = RawCompanyBalanceImporter(sheet)
                raw_data = imp.import_from_excel()
                completed = CompanyPreprocessorDAO.process_company_data(raw_data)
                imported.append(completed)
            except Exception as e:
                print("ERRORE: Impossibile importare la company dallo sheet " + sheet.title)
                print(e.with_traceback(None))
                print(e)

        wb.close()
        return imported

    @property
    def company(self) -> CompanyData:
        if self.__company__ is None:
            raise Exception("Importazione non ancora avviata")

        return self.__company__

    @property
    def translations(self) -> dict:
        if self.__translations__ is None:
            raise Exception("Importazione non ancora avviata")

        return self.__translations__